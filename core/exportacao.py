from __future__ import annotations
import io
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.styles import Font, PatternFill, Alignment
from copy import copy
from openpyxl.utils import get_column_letter

FORMATO_MOEDA = 'R$ #,##0.00'
FORMATO_INTEIRO = '0'
FORMATO_DECIMAL = '0.00'


def _coagir_numerico(df: pd.DataFrame, colunas) -> pd.DataFrame:
    out = df.copy()
    for coluna in colunas:
        if coluna not in out.columns:
            continue
        def conv(v):
            if pd.isna(v) or v == '':
                return float('nan')
            if isinstance(v, bool):
                return int(v)
            if isinstance(v, (int, float)):
                return float(v)
            s = str(v).strip().replace('\xa0', '').replace('R$', '').strip()
            if ',' in s:
                s = s.replace('.', '').replace(',', '.')
            return float(s)
        out[coluna] = out[coluna].map(conv).astype('float64')
    return out


def _ajustar_larguras(ws):
    for coluna in ws.columns:
        letra = get_column_letter(coluna[0].column)
        maior = max((len('' if c.value is None else str(c.value)) for c in coluna), default=0)
        ws.column_dimensions[letra].width = min(max(maior + 2, 11), 46)


def _cabecalho(ws):
    fill = PatternFill('solid', fgColor='17365D')
    for c in ws[1]:
        c.fill = fill
        c.font = Font(color='FFFFFF', bold=True)
        c.alignment = Alignment(horizontal='center', vertical='center')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    _ajustar_larguras(ws)


def _formatar_colunas_por_nome(ws, moeda=(), inteiro=(), decimal=()):
    cabecalhos = {cel.value: cel.column for cel in ws[1] if cel.value is not None}
    for nome in moeda:
        col = cabecalhos.get(nome)
        if not col: continue
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            if cell.value is not None:
                if isinstance(cell.value, str):
                    s = cell.value.strip().replace('\xa0', '').replace('R$', '').strip()
                    if ',' in s: s = s.replace('.', '').replace(',', '.')
                    cell.value = float(s)
                cell.number_format = FORMATO_MOEDA
    for nome in inteiro:
        col = cabecalhos.get(nome)
        if not col: continue
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row,column=col).value is not None:
                ws.cell(row=row,column=col).number_format = FORMATO_INTEIRO
    for nome in decimal:
        col = cabecalhos.get(nome)
        if not col: continue
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row,column=col).value is not None:
                ws.cell(row=row,column=col).number_format = FORMATO_DECIMAL
    _cabecalho(ws)


def _extrato_centro_custo(det: pd.DataFrame) -> pd.DataFrame:
    if det.empty:
        return pd.DataFrame(columns=['Centro de Custo','Quantidade_OS','Horas_Pagas','Valor_Horario_Comercial','Valor_Fora_Nao_Plantao','Valor_Plantao_FDS','Total'])
    x = det.copy()
    fora = x['Regra'].astype(str).str.contains('FORA', case=False, na=False)
    if 'Plantão FDS' in x.columns:
        plantao = x['Plantão FDS'].fillna(False).astype(bool)
    else:
        plantao = pd.Series(False, index=x.index)
    x['Valor_Horario_Comercial'] = x['Valor Total'].where(~fora, 0.0)
    x['Valor_Plantao_FDS'] = x['Valor Total'].where(plantao, 0.0)
    x['Valor_Fora_Nao_Plantao'] = x['Valor Total'].where(fora & ~plantao, 0.0)
    return x.groupby('Centro de Custo', as_index=False).agg(
        Quantidade_OS=('OS','count'),
        Horas_Pagas=('Horas Pagas','sum'),
        Valor_Horario_Comercial=('Valor_Horario_Comercial','sum'),
        Valor_Fora_Nao_Plantao=('Valor_Fora_Nao_Plantao','sum'),
        Valor_Plantao_FDS=('Valor_Plantao_FDS','sum'),
        Total=('Valor Total','sum'),
    ).sort_values('Total', ascending=False)


def _centro_x_tecnico(det: pd.DataFrame) -> pd.DataFrame:
    if det.empty:
        return pd.DataFrame(columns=['Centro de Custo'])
    pv = pd.pivot_table(det, index='Centro de Custo', columns='Técnico', values='Valor Total', aggfunc='sum', fill_value=0, margins=True, margins_name='TOTAL')
    return pv.reset_index()


def gerar_excel_apuracao(det: pd.DataFrame, pend: pd.DataFrame, param: pd.DataFrame | None = None) -> bytes:
    det_exp = _coagir_numerico(det, ['Valor 1ª Hora','Valor Hora Adicional','Valor Mão de Obra','Quilometragem','Pedágio','Outros Ajustes','Valor Total','Horas Pagas','Tempo minutos'])
    pend_exp = pend.copy()

    if len(det_exp):
        por_tecnico = det_exp.groupby('Técnico', as_index=False).agg(Quantidade_OS=('OS','count'),Horas_Pagas=('Horas Pagas','sum'),Total=('Valor Total','sum'))
        por_cc = det_exp.groupby('Centro de Custo', as_index=False).agg(Quantidade_OS=('OS','count'),Horas_Pagas=('Horas Pagas','sum'),Total=('Valor Total','sum'))
        resumo = pd.DataFrame([
            {'Indicador':'OS processadas','Valor':int(len(det_exp))},
            {'Indicador':'Pendências','Valor':int(len(pend_exp))},
            {'Indicador':'Técnicos','Valor':int(det_exp['Técnico'].nunique())},
            {'Indicador':'Total a pagar','Valor':float(det_exp['Valor Total'].sum())},
        ])
    else:
        por_tecnico = pd.DataFrame(columns=['Técnico','Quantidade_OS','Horas_Pagas','Total'])
        por_cc = pd.DataFrame(columns=['Centro de Custo','Quantidade_OS','Horas_Pagas','Total'])
        resumo = pd.DataFrame([{'Indicador':'OS processadas','Valor':0},{'Indicador':'Pendências','Valor':len(pend_exp)},{'Indicador':'Técnicos','Valor':0},{'Indicador':'Total a pagar','Valor':0.0}])

    extrato = _extrato_centro_custo(det_exp)
    matriz_cc_tecnico = _centro_x_tecnico(det_exp)
    plantao = det_exp[det_exp['Plantão FDS'].fillna(False).astype(bool)].copy() if 'Plantão FDS' in det_exp.columns else pd.DataFrame(columns=det_exp.columns)
    param_exp = _coagir_numerico(param, ['hc_primeira','hc_adicional','fora_primeira','fora_adicional']) if param is not None else None

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        resumo.to_excel(writer,index=False,sheet_name='RESUMO')
        extrato.to_excel(writer,index=False,sheet_name='EXTRATO CENTRO CUSTO',startrow=0)
        matriz_cc_tecnico.to_excel(writer,index=False,sheet_name='CC X TÉCNICO')
        det_exp.to_excel(writer,index=False,sheet_name='DETALHAMENTO OS')
        por_tecnico.to_excel(writer,index=False,sheet_name='POR TÉCNICO')
        por_cc.to_excel(writer,index=False,sheet_name='POR CENTRO CUSTO')
        plantao.to_excel(writer,index=False,sheet_name='PLANTÃO FDS')
        pend_exp.to_excel(writer,index=False,sheet_name='PENDÊNCIAS')
        if param_exp is not None:
            param_exp.to_excel(writer,index=False,sheet_name='PARAMETRIZAÇÃO')

    buf.seek(0)
    wb = load_workbook(buf)

    _formatar_colunas_por_nome(wb['DETALHAMENTO OS'], moeda=['Valor 1ª Hora','Valor Hora Adicional','Valor Mão de Obra','Quilometragem','Pedágio','Outros Ajustes','Valor Total'], inteiro=['Horas Pagas','Tempo minutos'])
    _formatar_colunas_por_nome(wb['POR TÉCNICO'], moeda=['Total'], inteiro=['Quantidade_OS','Horas_Pagas'])
    _formatar_colunas_por_nome(wb['POR CENTRO CUSTO'], moeda=['Total'], inteiro=['Quantidade_OS','Horas_Pagas'])
    _formatar_colunas_por_nome(wb['EXTRATO CENTRO CUSTO'], moeda=['Valor_Horario_Comercial','Valor_Fora_Nao_Plantao','Valor_Plantao_FDS','Total'], inteiro=['Quantidade_OS','Horas_Pagas'])

    ws = wb['CC X TÉCNICO']
    _cabecalho(ws)
    for row in range(2, ws.max_row+1):
        for col in range(2, ws.max_column+1):
            if isinstance(ws.cell(row,col).value,(int,float)):
                ws.cell(row,col).number_format=FORMATO_MOEDA

    if 'PLANTÃO FDS' in wb.sheetnames:
        _formatar_colunas_por_nome(wb['PLANTÃO FDS'], moeda=['Valor 1ª Hora','Valor Hora Adicional','Valor Total'], inteiro=['Horas Pagas','Tempo minutos'])

    ws=wb['RESUMO']; _cabecalho(ws)
    for row in range(2,ws.max_row+1):
        if ws.cell(row,1).value=='Total a pagar': ws.cell(row,2).number_format=FORMATO_MOEDA

    if 'PARAMETRIZAÇÃO' in wb.sheetnames:
        _formatar_colunas_por_nome(wb['PARAMETRIZAÇÃO'], moeda=['hc_primeira','hc_adicional','fora_primeira','fora_adicional'])
    _cabecalho(wb['PENDÊNCIAS'])

    wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True; wb.calculation.calcMode='auto'
    out=io.BytesIO(); wb.save(out); return out.getvalue()


def gerar_excel_revisao_tecnico(arquivo_original: bytes, det: pd.DataFrame, pend: pd.DataFrame, tecnico: str = 'Todos os técnicos', revisao_id: str = '', versao: int = 1) -> bytes:
    """Planilha compacta para conferência do técnico.

    O relatório do Inside é apenas a origem. A revisão expõe somente os campos
    necessários para reconhecer o atendimento. O técnico pode editar Tipo de
    Serviço, Valor Fechado, KM, Pedágio e Observação/Justificativa. Os resultados
    do Excel são apenas uma prévia; o sistema recalcula o retorno antes da aprovação.
    """
    import json, unicodedata, sqlite3
    from openpyxl import Workbook
    from openpyxl.styles import Protection
    from openpyxl.worksheet.datavalidation import DataValidation
    from .storage import caminho_banco
    from .tabela_valores import tarifa_oficial

    if not arquivo_original:
        raise ValueError('Arquivo original não disponível.')
    src = load_workbook(io.BytesIO(arquivo_original), data_only=False)
    wso = src[src.sheetnames[0]]

    def norm(v):
        t=''.join(c for c in unicodedata.normalize('NFKD',str(v or '')) if not unicodedata.combining(c))
        return ' '.join(t.upper().split())

    hmap={norm(wso.cell(1,c).value):c for c in range(1,wso.max_column+1)}
    # Recorte deliberadamente compacto: não replicar o relatório inteiro do Inside.
    desejadas=['Ordem','Status','Abertura','Cliente','Fantasia','Defeito Informado','Técnico','Fechamento','Defeito',
               'Observações Abertura','Observações Fechamento','DATA DE ATENDIMENTO','HORA INCIO','HORA FIM',
               'TEMPO DE ATENDIMENTO','SERVIÇO']
    editaveis=['Tipo de Serviço','Valor Fechado','KM','Pedágio','Peças / Materiais (R$)','Observação/Justificativa do Técnico']
    sistema=['SISTEMA - Tempo Calculado (min)','SISTEMA - Horas Pagas','SISTEMA - Regra Aplicada',
             'SISTEMA - Valor 1ª Hora','SISTEMA - Valor Hora Adicional','SISTEMA - Deslocamento',
             'SISTEMA - Total Previsto']
    headers=desejadas+editaveis+sistema+['SISTEMA - Linha Origem']

    detf=det.copy() if det is not None else pd.DataFrame()
    pendf=pend.copy() if pend is not None else pd.DataFrame()
    if tecnico!='Todos os técnicos':
        if len(detf): detf=detf[detf['Técnico'].astype(str)==str(tecnico)]
        if len(pendf): pendf=pendf[pendf['Técnico'].astype(str)==str(tecnico)]
    linhas={}
    for _,r in detf.iterrows(): linhas[int(r['linha_origem'])]=r.to_dict()
    for _,r in pendf.iterrows(): linhas.setdefault(int(r['linha_origem']),r.to_dict())
    if not linhas: raise ValueError('Não existem atendimentos para o técnico selecionado.')

    # Opções válidas: tabela ativa + tipos já presentes na apuração (retrocompatibilidade).
    tipos=set()
    try:
        db=caminho_banco()
        if db.exists():
            with sqlite3.connect(db) as con:
                tipos.update(str(r[0]).strip() for r in con.execute("SELECT descricao FROM tarifas WHERE ativo=1 AND descricao IS NOT NULL") if str(r[0]).strip())
    except Exception:
        pass
    for calc in linhas.values():
        for k in ('Tipo Pagamento','Tipo OS'):
            if str(calc.get(k,'') or '').strip(): tipos.add(str(calc.get(k)).strip())
    tipos=sorted(tipos, key=lambda x:norm(x))

    wb=Workbook(); ws=wb.active; ws.title='REVISÃO TÉCNICA'
    for c,h in enumerate(headers,1):
        x=ws.cell(1,c,h); x.font=Font(bold=True,color='FFFFFF'); x.fill=PatternFill('solid',fgColor='17365D'); x.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)

    # Aba oculta para lista suspensa e tarifas específicas por linha/técnico/regra.
    # RC14: primeira hora e adicional ficam em células NUMÉRICAS separadas.
    aux=wb.create_sheet('_SISTEMA_TARIFAS'); aux.sheet_state='hidden'
    aux.cell(1,1,'Tipo de Serviço')
    for j,li in enumerate(sorted(linhas),0):
        aux.cell(1,2+(j*2),f'V1_R{j+2}')
        aux.cell(1,3+(j*2),f'VA_R{j+2}')
    for i,t in enumerate(tipos,2): aux.cell(i,1,t)

    base={}
    linha_excel_por_origem={}
    for rr,li in enumerate(sorted(linhas),2):
        linha_excel_por_origem[li]=rr
        calc=linhas[li]
        for c,h in enumerate(desejadas,1):
            oc=hmap.get(norm(h)); ws.cell(rr,c,wso.cell(li,oc).value if oc else '')
        tipo=calc.get('Tipo Pagamento',calc.get('Tipo OS',''))
        mao=float(calc.get('Valor Mão de Obra',calc.get('Valor Total',0)) or 0)
        origem_valor=norm(calc.get('Origem do Valor',''))
        valor_fechado=mao if ('MANUAL' in origem_valor or 'VALOR FECHADO' in norm(tipo)) else ''
        vals={'Tipo de Serviço':tipo,'Valor Fechado':valor_fechado,'KM':calc.get('Quilometragem',0),'Pedágio':calc.get('Pedágio',0),
              'Peças / Materiais (R$)':calc.get('Peças / Materiais (R$)',calc.get('Peças/Materiais',0)) or 0,
              'Observação/Justificativa do Técnico':'','SISTEMA - Tempo Calculado (min)':calc.get('Tempo minutos'),
              'SISTEMA - Horas Pagas':calc.get('Horas Pagas'),'SISTEMA - Regra Aplicada':calc.get('Regra',''),
              'SISTEMA - Valor 1ª Hora':calc.get('Valor 1ª Hora'),'SISTEMA - Valor Hora Adicional':calc.get('Valor Hora Adicional'),
              'SISTEMA - Deslocamento':0,'SISTEMA - Total Previsto':0,'SISTEMA - Linha Origem':li}
        for h,v in vals.items(): ws.cell(rr,headers.index(h)+1,v)
        base[str(li)]={'Tipo de Serviço':tipo,'Valor Fechado':valor_fechado,'KM':calc.get('Quilometragem',0),'Pedágio':calc.get('Pedágio',0),'Peças / Materiais (R$)':calc.get('Peças / Materiais (R$)',calc.get('Peças/Materiais',0)) or 0}
        ws.row_dimensions[rr].height=22

    # Tarifas auxiliares por OS. Cada OS tem duas colunas numéricas: V1 e VA.
    # A resolução usa o mesmo catálogo do sistema, incluindo aliases equivalentes.
    for j,li in enumerate(sorted(linhas),0):
        calc=linhas[li]; fora='FORA' in norm(calc.get('Regra',''))
        c_v1=2+(j*2); c_va=3+(j*2)
        for i,t in enumerate(tipos,2):
            tar=tarifa_oficial(tecnico,t,fora_horario=fora)
            if tar is not None:
                aux.cell(i,c_v1,float(tar[0]))
                aux.cell(i,c_va,float(tar[1]))

    col={h:headers.index(h)+1 for h in headers}
    # Lista suspensa real para Tipo de Serviço.
    if tipos:
        dv=DataValidation(type='list',formula1=f"'_SISTEMA_TARIFAS'!$A$2:$A${len(tipos)+1}",allow_blank=True)
        dv.error='Selecione um Tipo de Serviço disponível na lista.'; dv.errorTitle='Tipo inválido'; dv.prompt='Selecione o serviço usado para pagamento.'; dv.promptTitle='Tipo de Serviço'; dv.showInputMessage=True
        ws.add_data_validation(dv); dv.add(f"{get_column_letter(col['Tipo de Serviço'])}2:{get_column_letter(col['Tipo de Serviço'])}{ws.max_row}")

    # Fórmulas de prévia. O sistema continuará recalculando tudo na importação.
    for rr,li in enumerate(sorted(linhas),2):
        t=get_column_letter(col['Tipo de Serviço']); vf=get_column_letter(col['Valor Fechado']); km=get_column_letter(col['KM']); ped=get_column_letter(col['Pedágio']); pec=get_column_letter(col['Peças / Materiais (R$)'])
        hp=get_column_letter(col['SISTEMA - Horas Pagas']); v1=get_column_letter(col['SISTEMA - Valor 1ª Hora']); va=get_column_letter(col['SISTEMA - Valor Hora Adicional']); des=get_column_letter(col['SISTEMA - Deslocamento']); tot=get_column_letter(col['SISTEMA - Total Previsto'])
        idx=rr-2
        aux_v1=get_column_letter(2+(idx*2)); aux_va=get_column_letter(3+(idx*2))
        if tipos:
            ws[f'{v1}{rr}']=f'=IFERROR(INDEX(\'_SISTEMA_TARIFAS\'!${aux_v1}$2:${aux_v1}${len(tipos)+1},MATCH({t}{rr},\'_SISTEMA_TARIFAS\'!$A$2:$A${len(tipos)+1},0)),0)'
            ws[f'{va}{rr}']=f'=IFERROR(INDEX(\'_SISTEMA_TARIFAS\'!${aux_va}$2:${aux_va}${len(tipos)+1},MATCH({t}{rr},\'_SISTEMA_TARIFAS\'!$A$2:$A${len(tipos)+1},0)),0)'
        else:
            ws[f'{v1}{rr}']='=0'
            ws[f'{va}{rr}']='=0'
        ws[f'{des}{rr}']=f'=IFERROR({km}{rr}*1.5,0)'
        mao=f'IF(ISNUMBER(SEARCH("VALOR FECHADO",{t}{rr})),IFERROR({vf}{rr},0),{v1}{rr}+MAX({hp}{rr}-1,0)*{va}{rr})'
        ws[f'{tot}{rr}']=f'=ROUND({mao}+{des}{rr}+IFERROR({ped}{rr},0)+IFERROR({pec}{rr},0),2)'

    edit_fill=PatternFill('solid',fgColor='FFF2CC')
    for c,h in enumerate(headers,1):
        ws.column_dimensions[get_column_letter(c)].width=44 if 'Observações' in h or 'Justificativa' in h else min(max(len(h)+2,12),26)
        for r in range(2,ws.max_row+1):
            cell=ws.cell(r,c); cell.protection=Protection(locked=h not in editaveis)
            if h in editaveis: cell.fill=edit_fill
            if 'Observações' in h or 'Justificativa' in h: cell.alignment=Alignment(vertical='top',wrap_text=True)
            if h in ('Valor Fechado','Pedágio','Peças / Materiais (R$)','SISTEMA - Valor 1ª Hora','SISTEMA - Valor Hora Adicional','SISTEMA - Deslocamento','SISTEMA - Total Previsto'): cell.number_format=FORMATO_MOEDA
            if h=='KM': cell.number_format='0.0'
    ws.column_dimensions[get_column_letter(len(headers))].hidden=True
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions; ws.row_dimensions[1].height=30
    ws.sheet_view.topLeftCell='A1'; ws.sheet_view.selection[0].activeCell='A1'; ws.sheet_view.selection[0].sqref='A1'
    # Permite selecionar e editar as células desbloqueadas; mantém a base protegida.
    ws.protection.sheet=True; ws.protection.password='prosecurity'; ws.protection.selectLockedCells=True; ws.protection.selectUnlockedCells=False

    meta=wb.create_sheet('_SISTEMA_REVISAO'); meta.sheet_state='veryHidden'
    for r,(k,v) in enumerate([('revisao_id',revisao_id),('tecnico',tecnico),('versao',versao),('base_json',json.dumps(base,ensure_ascii=False))],1): meta.cell(r,1,k); meta.cell(r,2,str(v))
    wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True; wb.calculation.calcMode='auto'
    out=io.BytesIO(); wb.save(out); return out.getvalue()
