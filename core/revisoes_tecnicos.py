from __future__ import annotations
import io, json, uuid
from datetime import datetime
import pandas as pd
from openpyxl import load_workbook
from .historico import conectar, inicializar_banco

STATUS_ABERTA='AGUARDANDO RETORNO'

def inicializar_revisoes(db_path=None):
    inicializar_banco(db_path)
    with conectar(db_path) as con:
        con.executescript('''
        CREATE TABLE IF NOT EXISTS revisoes_tecnicos(
          id TEXT PRIMARY KEY, apuracao_id TEXT NOT NULL, tecnico TEXT NOT NULL, competencia TEXT,
          status TEXT NOT NULL, versao INTEGER NOT NULL DEFAULT 1, criado_em TEXT NOT NULL,
          recebido_em TEXT, concluido_em TEXT, aprovado_por TEXT, gerado_em TEXT
        );
        CREATE TABLE IF NOT EXISTS revisoes_itens(
          id INTEGER PRIMARY KEY AUTOINCREMENT, revisao_id TEXT NOT NULL, linha_origem INTEGER NOT NULL,
          os TEXT, campo TEXT NOT NULL, valor_base TEXT, valor_informado TEXT, justificativa TEXT,
          status TEXT NOT NULL DEFAULT 'PENDENTE', motivo_rejeicao TEXT, versao INTEGER NOT NULL DEFAULT 1,
          decidido_por TEXT, decidido_em TEXT
        );
        CREATE TABLE IF NOT EXISTS revisoes_arquivos(
          id INTEGER PRIMARY KEY AUTOINCREMENT, revisao_id TEXT NOT NULL, versao INTEGER NOT NULL,
          tipo TEXT NOT NULL, nome_arquivo TEXT, recebido_em TEXT NOT NULL
        );
        ''')
        cols={r[1] for r in con.execute("PRAGMA table_info(revisoes_tecnicos)").fetchall()}
        if 'gerado_em' not in cols:
            con.execute("ALTER TABLE revisoes_tecnicos ADD COLUMN gerado_em TEXT")

def obter_ou_criar_revisao(apuracao_id, tecnico, competencia='', db_path=None):
    inicializar_revisoes(db_path)
    with conectar(db_path) as con:
        r=con.execute("SELECT * FROM revisoes_tecnicos WHERE apuracao_id=? AND tecnico=? ORDER BY criado_em DESC LIMIT 1",(apuracao_id,tecnico)).fetchone()
        if r: return dict(r)
        rid=str(uuid.uuid4()); agora=datetime.now().isoformat(timespec='seconds')
        con.execute("INSERT INTO revisoes_tecnicos(id,apuracao_id,tecnico,competencia,status,versao,criado_em) VALUES(?,?,?,?,?,1,?)",(rid,apuracao_id,tecnico,competencia,STATUS_ABERTA,agora))
        return {'id':rid,'apuracao_id':apuracao_id,'tecnico':tecnico,'competencia':competencia,'status':STATUS_ABERTA,'versao':1,'criado_em':agora}

def iniciar_revisao(revisao_id, db_path=None):
    inicializar_revisoes(db_path)
    agora=datetime.now().isoformat(timespec='seconds')
    with conectar(db_path) as con:
        con.execute("UPDATE revisoes_tecnicos SET gerado_em=COALESCE(gerado_em,?), status=? WHERE id=?",(agora,STATUS_ABERTA,revisao_id))
    return agora

def listar_revisoes(apuracao_id, db_path=None):
    inicializar_revisoes(db_path)
    with conectar(db_path) as con:
        return pd.read_sql_query('''SELECT r.*, 
          SUM(CASE WHEN i.status='PENDENTE' THEN 1 ELSE 0 END) pendentes,
          SUM(CASE WHEN i.status='APROVADO' THEN 1 ELSE 0 END) aprovados,
          SUM(CASE WHEN i.status='REJEITADO' THEN 1 ELSE 0 END) rejeitados
          FROM revisoes_tecnicos r LEFT JOIN revisoes_itens i ON i.revisao_id=r.id
          WHERE r.apuracao_id=? AND r.gerado_em IS NOT NULL GROUP BY r.id ORDER BY r.tecnico''',con,params=(apuracao_id,))

def itens_revisao(revisao_id, db_path=None):
    inicializar_revisoes(db_path)
    with conectar(db_path) as con:
        return pd.read_sql_query('SELECT * FROM revisoes_itens WHERE revisao_id=? ORDER BY linha_origem,campo,id',con,params=(revisao_id,))

def registrar_retorno(revisao_id, arquivo: bytes, nome_arquivo='', db_path=None):
    inicializar_revisoes(db_path)
    wb=load_workbook(io.BytesIO(arquivo),data_only=False)
    if '_SISTEMA_REVISAO' not in wb.sheetnames: raise ValueError('Arquivo não pertence ao fluxo de Revisão Técnica.')
    meta=wb['_SISTEMA_REVISAO']
    md={str(meta.cell(r,1).value):str(meta.cell(r,2).value or '') for r in range(1,meta.max_row+1)}
    if md.get('revisao_id') != str(revisao_id): raise ValueError('Esta planilha pertence a outra revisão/técnico.')
    ws=wb[wb.sheetnames[0]]
    headers={str(ws.cell(1,c).value or '').strip():c for c in range(1,ws.max_column+1)}
    # RC10 aceita o novo rótulo e mantém retrocompatibilidade com planilhas antigas.
    tipo_col='Tipo de Serviço' if 'Tipo de Serviço' in headers else 'Tipo para Pagamento'
    obrig=['Ordem',tipo_col,'Valor Fechado','KM','Pedágio','Observação/Justificativa do Técnico','SISTEMA - Linha Origem']
    pecas_col='Peças / Materiais (R$)' if 'Peças / Materiais (R$)' in headers else ('Peças/Materiais' if 'Peças/Materiais' in headers else None)
    falt=[x for x in obrig if x not in headers]
    if falt: raise ValueError('Planilha inválida. Colunas ausentes: '+', '.join(falt))
    base=json.loads(md.get('base_json') or '{}')
    versao=int(md.get('versao') or 1)
    itens=[]
    for r in range(2,ws.max_row+1):
        li=str(ws.cell(r,headers['SISTEMA - Linha Origem']).value or '').strip()
        if not li: continue
        b=base.get(li,{})
        just=str(ws.cell(r,headers['Observação/Justificativa do Técnico']).value or '').strip()
        for campo in ['Tipo de Serviço','Valor Fechado','KM','Pedágio'] + (['Peças / Materiais (R$)'] if pecas_col else []):
            coluna=tipo_col if campo=='Tipo de Serviço' else (pecas_col if campo=='Peças / Materiais (R$)' else campo)
            novo=ws.cell(r,headers[coluna]).value
            antigo=b.get(campo, b.get('Tipo para Pagamento','') if campo=='Tipo de Serviço' else '')
            def norm(v):
                if v is None: return ''
                if campo in ('Valor Fechado','KM','Pedágio','Peças / Materiais (R$)'):
                    try:return str(round(float(v),2))
                    except:return str(v).strip()
                return str(v).strip()
            if norm(novo)!=norm(antigo): itens.append((int(float(li)),str(ws.cell(r,headers['Ordem']).value or ''),campo,str(antigo),str(novo or ''),just))
    with conectar(db_path) as con:
        con.execute("DELETE FROM revisoes_itens WHERE revisao_id=? AND status='PENDENTE'",(revisao_id,))
        for li,osv,campo,antigo,novo,just in itens:
            con.execute("UPDATE revisoes_itens SET status='SUPERADO' WHERE revisao_id=? AND linha_origem=? AND campo=? AND status='REJEITADO'",(revisao_id,li,campo))
            con.execute('''INSERT INTO revisoes_itens(revisao_id,linha_origem,os,campo,valor_base,valor_informado,justificativa,status,versao)
                           VALUES(?,?,?,?,?,?,?,'PENDENTE',?)''',(revisao_id,li,osv,campo,antigo,novo,just,versao))
        agora=datetime.now().isoformat(timespec='seconds')
        con.execute("UPDATE revisoes_tecnicos SET status='EM CONFERÊNCIA', recebido_em=? WHERE id=?",(agora,revisao_id))
        con.execute("INSERT INTO revisoes_arquivos(revisao_id,versao,tipo,nome_arquivo,recebido_em) VALUES(?,?, 'RETORNO', ?,?)",(revisao_id,versao,nome_arquivo,agora))
    return len(itens)

def decidir_itens(revisao_id, ids, aprovar: bool, usuario, motivo='', db_path=None):
    if not ids:return 0
    status='APROVADO' if aprovar else 'REJEITADO'; agora=datetime.now().isoformat(timespec='seconds')
    qs=','.join('?' for _ in ids)
    with conectar(db_path) as con:
        con.execute(f"UPDATE revisoes_itens SET status=?,motivo_rejeicao=?,decidido_por=?,decidido_em=? WHERE revisao_id=? AND id IN ({qs})",(status,motivo,usuario,agora,revisao_id,*ids))
        pend=con.execute("SELECT COUNT(*) FROM revisoes_itens WHERE revisao_id=? AND status='PENDENTE'",(revisao_id,)).fetchone()[0]
        rej=con.execute("SELECT COUNT(*) FROM revisoes_itens WHERE revisao_id=? AND status='REJEITADO'",(revisao_id,)).fetchone()[0]
        novo='AGUARDANDO CORREÇÃO' if not pend and rej else ('REVISÃO CONCLUÍDA' if not pend else 'EM CONFERÊNCIA')
        con.execute("UPDATE revisoes_tecnicos SET status=? WHERE id=?",(novo,revisao_id))
    return len(ids)

def concluir_revisao(revisao_id, usuario, db_path=None):
    with conectar(db_path) as con:
        pend=con.execute("SELECT COUNT(*) FROM revisoes_itens WHERE revisao_id=? AND status NOT IN ('APROVADO','SUPERADO')",(revisao_id,)).fetchone()[0]
        if pend: raise ValueError('Ainda existem itens pendentes ou rejeitados nesta revisão.')
        con.execute("UPDATE revisoes_tecnicos SET status='APROVADA',concluido_em=?,aprovado_por=? WHERE id=?",(datetime.now().isoformat(timespec='seconds'),usuario,revisao_id))

def preparar_planilha_correcao(arquivo_revisao: bytes, itens_rejeitados: pd.DataFrame) -> bytes:
    """Reduz a planilha de revisão às OS devolvidas e inclui o motivo da correção."""
    wb=load_workbook(io.BytesIO(arquivo_revisao))
    ws=wb[wb.sheetnames[0]]
    headers={str(ws.cell(1,c).value or '').strip():c for c in range(1,ws.max_column+1)}
    cli=headers.get('SISTEMA - Linha Origem')
    if not cli: raise ValueError('Identificador interno da revisão ausente.')
    motivos={int(r['linha_origem']):str(r.get('motivo_rejeicao') or '') for _,r in itens_rejeitados.iterrows()}
    cm=ws.max_column+1; ws.cell(1,cm,'Motivo da Correção')
    ws.cell(1,cm).font=ws.cell(1,1).font.copy(); ws.cell(1,cm).fill=ws.cell(1,1).fill.copy()
    for r in range(ws.max_row,1,-1):
        try: li=int(float(ws.cell(r,cli).value))
        except: li=-1
        if li not in motivos: ws.delete_rows(r,1)
        else: ws.cell(r,cm,motivos[li])
    meta=wb['_SISTEMA_REVISAO']; md={str(meta.cell(r,1).value):r for r in range(1,meta.max_row+1)}
    rv=md.get('versao'); atual=int(meta.cell(rv,2).value or 1) if rv else 1
    if rv: meta.cell(rv,2,atual+1)
    out=io.BytesIO(); wb.save(out); return out.getvalue()
