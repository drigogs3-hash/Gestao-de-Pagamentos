import io
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.filters import FilterColumn, CustomFilters, CustomFilter

from core.exportacao import gerar_excel_revisao_tecnico


def _original_problematico():
    wb=Workbook(); ws=wb.active; ws.title='Relatorio'
    ws.append(['Cliente','Ordem','Técnico'])
    for r in range(2,301):
        ws.append([f'Cliente {r}', r, 'RUBENS' if r in (253,260) else 'OUTRO'])
    for r in range(2,253):
        ws.row_dimensions[r].hidden=True
        ws.row_dimensions[r].outlineLevel=1
        ws.row_dimensions[r].collapsed=True
    ws.auto_filter.ref='A1:C300'
    fc=FilterColumn(colId=2, customFilters=CustomFilters(customFilter=[CustomFilter(operator='equal', val='RUBENS')]))
    ws.auto_filter.filterColumn.append(fc)
    ws.sheet_view.topLeftCell='A253'
    ws.sheet_view.selection[0].activeCell='A253'; ws.sheet_view.selection[0].sqref='A253'
    b=io.BytesIO(); wb.save(b); return b.getvalue()


def test_revisao_abre_na_linha_1_sem_ocultacao_ou_filtro_herdado():
    det=pd.DataFrame([
        {'linha_origem':253,'Técnico':'RUBENS','OS':'253','Cliente':'Cliente 253','Tempo minutos':60,'Horas Pagas':1,'Regra':'HORÁRIO COMERCIAL','Valor 1ª Hora':200,'Valor Hora Adicional':90,'Valor Total':200,'Tipo OS':'PORTÃO','Centro de Custo':'PRO'},
        {'linha_origem':260,'Técnico':'RUBENS','OS':'260','Cliente':'Cliente 260','Tempo minutos':60,'Horas Pagas':1,'Regra':'HORÁRIO COMERCIAL','Valor 1ª Hora':200,'Valor Hora Adicional':90,'Valor Total':200,'Tipo OS':'PORTÃO','Centro de Custo':'PRO'},
    ])
    out=gerar_excel_revisao_tecnico(_original_problematico(),det,pd.DataFrame(), 'RUBENS')
    wb=load_workbook(io.BytesIO(out)); ws=wb[wb.sheetnames[0]]
    assert ws.max_row == 3
    assert [ws.cell(r,1).value for r in range(1,4)] == ['Ordem',253,260]
    assert all(not ws.row_dimensions[r].hidden for r in range(1,ws.max_row+1))
    assert all(ws.row_dimensions[r].outlineLevel == 0 for r in range(1,ws.max_row+1))
    assert all(not ws.row_dimensions[r].collapsed for r in range(1,ws.max_row+1))
    assert len(ws.auto_filter.filterColumn) == 0
    assert ws.sheet_view.topLeftCell == 'A1'
    assert ws.sheet_view.selection[0].activeCell == 'A1'
    assert str(ws.sheet_view.selection[0].sqref) == 'A1'


def test_revisao_limita_altura_das_observacoes_sem_perder_texto():
    wb=Workbook(); ws=wb.active; ws.title='Relatorio'
    ws.append(['Cliente','Ordem','Técnico','Observações   Abertura','Observações   Fechamento'])
    abertura='Abertura muito longa ' * 40
    fechamento='Fechamento muito longo ' * 60
    ws.append(['Cliente 2',2,'RUBENS',abertura,fechamento])
    ws.row_dimensions[2].height=240
    b=io.BytesIO(); wb.save(b)
    det=pd.DataFrame([{'linha_origem':2,'Técnico':'RUBENS','OS':'2','Cliente':'Cliente 2','Tempo minutos':60,'Horas Pagas':1,'Regra':'HORÁRIO COMERCIAL','Valor 1ª Hora':200,'Valor Hora Adicional':90,'Valor Total':200,'Tipo OS':'PORTÃO','Centro de Custo':'PRO'}])
    out=gerar_excel_revisao_tecnico(b.getvalue(),det,pd.DataFrame(),'RUBENS')
    wb2=load_workbook(io.BytesIO(out)); ws2=wb2[wb2.sheetnames[0]]
    hs={ws2.cell(1,c).value:c for c in range(1,ws2.max_column+1)}
    ca=hs['Observações Abertura']; cf=hs['Observações Fechamento']
    assert ws2.row_dimensions[2].height == 22
    assert ws2.cell(2,ca).value == abertura
    assert ws2.cell(2,cf).value == fechamento
    assert ws2.cell(2,ca).alignment.wrap_text is True
    assert ws2.cell(2,cf).alignment.wrap_text is True
    from openpyxl.utils import get_column_letter
    assert ws2.column_dimensions[get_column_letter(ca)].width == 44
    assert ws2.column_dimensions[get_column_letter(cf)].width == 44


def test_rc14_tarifas_da_revisao_sao_numericas_e_formula_nao_usa_vlookup_textual():
    wb=Workbook(); ws=wb.active; ws.title='Relatorio'
    ws.append(['Cliente','Ordem','Técnico','SERVIÇO'])
    ws.append(['Cliente A',1,'TESTE','ALARME'])
    b=io.BytesIO(); wb.save(b)
    det=pd.DataFrame([{'linha_origem':2,'Técnico':'TESTE','OS':'1','Cliente':'Cliente A','Tempo minutos':60,'Horas Pagas':1,'Regra':'HORÁRIO COMERCIAL','Valor 1ª Hora':90.10,'Valor Hora Adicional':53.0,'Valor Total':90.10,'Tipo OS':'ALARME','Tipo Pagamento':'ALARME','Centro de Custo':'PRO'}])
    out=gerar_excel_revisao_tecnico(b.getvalue(),det,pd.DataFrame(),'TESTE')
    w=load_workbook(io.BytesIO(out),data_only=False); s=w['REVISÃO TÉCNICA']; aux=w['_SISTEMA_TARIFAS']
    hs={s.cell(1,c).value:c for c in range(1,s.max_column+1)}
    assert aux['A2'].value == 'ALARME'
    assert aux['B2'].value == 90.10
    assert aux['C2'].value == 53.0
    f1=s.cell(2,hs['SISTEMA - Valor 1ª Hora']).value
    ft=s.cell(2,hs['SISTEMA - Total Previsto']).value
    assert 'INDEX(' in f1 and 'MATCH(' in f1 and 'VLOOKUP' not in f1
    assert 'ROUND(' in ft
