import io, tempfile
import pandas as pd
from openpyxl import Workbook, load_workbook
from core.exportacao import gerar_excel_revisao_tecnico
from core.revisoes_tecnicos import inicializar_revisoes, obter_ou_criar_revisao, registrar_retorno, itens_revisao, decidir_itens, preparar_planilha_correcao

def original():
    wb=Workbook(); ws=wb.active
    hs=['Ordem','Cliente','Técnico','Defeito','Observações Abertura','Observações Fechamento','SERVIÇO']
    ws.append(hs); ws.append([100,'Cliente A','ALEX','CONTROLE ACESSO','texto '*100,'fechamento '*100,'(MAN) CONTROLE'])
    b=io.BytesIO(); wb.save(b); return b.getvalue()

def det():
    return pd.DataFrame([{'linha_origem':2,'OS':'100','Cliente':'Cliente A','Técnico':'ALEX','Tipo OS':'CONTROLE ACESSO','Tipo Pagamento':'CONTROLE ACESSO','Tempo minutos':60,'Horas Pagas':1,'Regra':'HC','Valor 1ª Hora':127.2,'Valor Hora Adicional':53,'Valor Total':127.2,'Quilometragem':0,'Pedágio':0}])

def test_exportacao_so_campos_necessarios_e_protegida():
    data=gerar_excel_revisao_tecnico(original(),det(),pd.DataFrame(),'ALEX','rev-1',1)
    wb=load_workbook(io.BytesIO(data)); ws=wb[wb.sheetnames[0]]
    headers=[ws.cell(1,c).value for c in range(1,ws.max_column+1)]
    assert 'Tipo de Serviço' in headers and 'Valor Fechado' in headers and 'KM' in headers and 'Pedágio' in headers
    assert 'SISTEMA - Plantão FDS' not in headers
    assert ws.protection.sheet
    assert not ws.cell(2,headers.index('KM')+1).protection.locked
    assert ws.cell(2,headers.index('Ordem')+1).protection.locked
    assert ws.row_dimensions[2].height==22
    assert wb['_SISTEMA_REVISAO'].sheet_state=='veryHidden'

def test_retorno_isolado_por_revisao_e_decisao_em_massa():
    with tempfile.NamedTemporaryFile(suffix='.db') as f:
        inicializar_revisoes(f.name)
        r=obter_ou_criar_revisao('ap1','ALEX','09/2026',f.name)
        data=gerar_excel_revisao_tecnico(original(),det(),pd.DataFrame(),'ALEX',r['id'],1)
        wb=load_workbook(io.BytesIO(data)); ws=wb[wb.sheetnames[0]]
        hs={ws.cell(1,c).value:c for c in range(1,ws.max_column+1)}
        ws.protection.sheet=False; ws.cell(2,hs['KM'],42); ws.cell(2,hs['Pedágio'],18.4); ws.cell(2,hs['Observação/Justificativa do Técnico'],'visita externa')
        b=io.BytesIO(); wb.save(b)
        assert registrar_retorno(r['id'],b.getvalue(),'ret.xlsx',f.name)==2
        itens=itens_revisao(r['id'],f.name); assert set(itens.campo)=={'KM','Pedágio'}
        decidir_itens(r['id'],itens.id.astype(int).tolist(),True,'Rodrigo',db_path=f.name)
        assert set(itens_revisao(r['id'],f.name).status)=={'APROVADO'}

def test_planilha_correcao_contem_so_os_rejeitada():
    data=gerar_excel_revisao_tecnico(original(),det(),pd.DataFrame(),'ALEX','rev-1',1)
    rej=pd.DataFrame([{'linha_origem':2,'motivo_rejeicao':'Comprovar pedágio'}])
    corr=preparar_planilha_correcao(data,rej); wb=load_workbook(io.BytesIO(corr)); ws=wb[wb.sheetnames[0]]
    assert ws.max_row==2 and ws.cell(1,ws.max_column).value=='Motivo da Correção'
    assert ws.cell(2,ws.max_column).value=='Comprovar pedágio'

def test_rc10_layout_compacto_campos_editaveis_e_previa():
    data=gerar_excel_revisao_tecnico(original(),det(),pd.DataFrame(),'ALEX','rev-rc10',1)
    wb=load_workbook(io.BytesIO(data),data_only=False); ws=wb['REVISÃO TÉCNICA']
    headers=[ws.cell(1,c).value for c in range(1,ws.max_column+1)]
    esperadas=['Ordem','Status','Abertura','Cliente','Fantasia','Defeito Informado','Técnico','Fechamento','Defeito','Observações Abertura','Observações Fechamento','DATA DE ATENDIMENTO','HORA INCIO','HORA FIM','TEMPO DE ATENDIMENTO','SERVIÇO','Tipo de Serviço','Valor Fechado','KM','Pedágio','Peças / Materiais (R$)','Observação/Justificativa do Técnico','SISTEMA - Tempo Calculado (min)','SISTEMA - Horas Pagas','SISTEMA - Regra Aplicada','SISTEMA - Valor 1ª Hora','SISTEMA - Valor Hora Adicional','SISTEMA - Deslocamento','SISTEMA - Total Previsto','SISTEMA - Linha Origem']
    assert headers == esperadas
    for h in ['Tipo de Serviço','Valor Fechado','KM','Pedágio','Peças / Materiais (R$)','Observação/Justificativa do Técnico']:
        assert ws.cell(2,headers.index(h)+1).protection.locked is False
    assert ws.cell(2,headers.index('Ordem')+1).protection.locked is True
    assert ws.data_validations.count == 1
    assert str(ws.cell(2,headers.index('SISTEMA - Deslocamento')+1).value).startswith('=IFERROR(')
    assert 'SEARCH("VALOR FECHADO"' in str(ws.cell(2,headers.index('SISTEMA - Total Previsto')+1).value)
    assert 'U2' in str(ws.cell(2,headers.index('SISTEMA - Total Previsto')+1).value)  # peças/materiais no total
    assert wb.calculation.calcMode == 'auto' and wb.calculation.fullCalcOnLoad
